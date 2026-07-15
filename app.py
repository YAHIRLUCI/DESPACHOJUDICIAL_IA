import os
import io
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import requests
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
import fitz
import pytesseract
from PIL import Image

app = Flask(__name__)
CORS(app)

# Ajusta esta ruta si usas Tesseract OCR en Windows
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
DOCS_DIR = './documentos'
if not os.path.exists(DOCS_DIR): os.makedirs(DOCS_DIR)

def extraer_texto(ruta):
    """Lee el documento de forma segura según su extensión."""
    try:
        if ruta.lower().endswith('.pdf'):
            doc = fitz.open(ruta)
            texto = "\n".join([p.get_text() for p in doc])
            return texto
        elif ruta.lower().endswith('.docx'):
            doc = Document(ruta)
            texto = "\n".join([para.text for para in doc.paragraphs])
            return texto
        else:
            with open(ruta, 'r', encoding='utf-8', errors='ignore') as f: 
                return f.read()
    except Exception as e:
        print(f"Error leyendo {ruta}: {e}")
        return ""

@app.route('/subir', methods=['POST'])
def subir():
    file = request.files['file']
    file.save(os.path.join(DOCS_DIR, file.filename))
    return jsonify({"mensaje": "OK"})

@app.route('/buscar', methods=['GET'])
def buscar():
    q = request.args.get('q', '').lower()
    resultados = []
    
    for f in os.listdir(DOCS_DIR):
        ruta = os.path.join(DOCS_DIR, f)
        texto_original = extraer_texto(ruta)
        texto_lower = texto_original.lower()
        
        if q in f.lower() or q in texto_lower:
            texto_limpio = texto_original.replace('\n', ' ').strip()
            extracto = texto_limpio[:150] + "..." if texto_limpio else "Documento sin texto extraíble."
                
            resultados.append({
                "nombre": f,
                "ruta": f,
                "extracto": extracto
            })
            
    return jsonify(resultados)

@app.route('/sinopsis_ia', methods=['POST'])
def sinopsis_ia():
    ruta_archivo = os.path.join(DOCS_DIR, request.json.get('ruta'))
    texto = extraer_texto(ruta_archivo)
    
    if not texto.strip():
        return jsonify({"sinopsis": "No se pudo extraer texto suficiente de este documento."})
    
    prompt = f"Eres un asistente legal. Escribe una sinopsis muy breve (máximo 3 líneas) explicando de qué trata exclusivamente este documento:\n\n{texto[:2500]}"
    try:
        res = requests.post("http://localhost:1234/v1/chat/completions", json={
            "messages": [{"role": "user", "content": prompt}],
            "temperature": 0.3, "max_tokens": 150
        })
        return jsonify({"sinopsis": res.json()['choices'][0]['message']['content']})
    except Exception as e: 
        return jsonify({"sinopsis": f"Error de conexión con LM Studio: {str(e)}"})

@app.route('/leer', methods=['GET'])
def leer():
    ruta = os.path.join(DOCS_DIR, request.args.get('ruta'))
    texto = extraer_texto(ruta)
    if texto.strip(): 
        return jsonify({"contenido": texto})
    return jsonify({"error": "No se pudo extraer texto o el archivo está vacío."})

@app.route('/resumir', methods=['POST'])
def resumir():
    texto = request.json.get('texto', '')
    prompt = f"Actúa como abogado experto. Realiza un resumen ejecutivo profesional y estructurado de este texto legal:\n\n{texto[:4000]}"
    try:
        res = requests.post("http://localhost:1234/v1/chat/completions", json={
            "messages": [{"role": "user", "content": prompt}],
            "temperature": 0.2, "max_tokens": 800
        })
        return jsonify({"resumen": res.json()['choices'][0]['message']['content']})
    except Exception as e: 
        return jsonify({"resumen": f"Error: {str(e)}"})

@app.route('/exportar', methods=['POST'])
def exportar():
    data = request.json
    cont = data.get('contenido', '')
    form = data.get('formato', 'docx')
    
    if form == 'pdf':
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer, pagesize=letter)
        width, height = letter
        c.setFont("Helvetica", 12)
        y = height - 50
        for line in cont.split('\n'):
            if y < 50: c.showPage(); y = height - 50
            c.drawString(50, y, line[:90])
            y -= 20
        c.save()
        buffer.seek(0)
        return send_file(buffer, mimetype='application/pdf', as_attachment=True, download_name='Resumen_Legal.pdf')
    else:
        doc = Document()
        doc.add_heading('Resumen Legal', 0)
        doc.add_paragraph(cont)
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return send_file(buffer, mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document', as_attachment=True, download_name='Resumen_Legal.docx')

if __name__ == '__main__': 
    app.run(port=5000)